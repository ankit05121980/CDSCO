#!/usr/bin/env python3
"""Build four DDRS knowledge-transfer documents (.docx):
  1) Data Migration from existing CDSCO portals to DDRS
  2) Transition / Refactoring / Migration / Upgradation & Maintenance (post Go-Live)
  3) Advanced Analytics & Intelligent Automation - Project Citation Format
  4) Intelligent Systems Architecture & Regulatory Use Cases
"""
import os, datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(__file__)
IMG = os.path.join(HERE, 'images')
NAVY = RGBColor(0x0b, 0x3d, 0x7b); SAFFRON = RGBColor(0xc0, 0x60, 0x10)
GREEN = RGBColor(0x0e, 0x66, 0x06); GREYC = RGBColor(0x55, 0x60, 0x6b)
TODAY = datetime.date.today().strftime('%d %B %Y')


class DDoc:
    def __init__(self):
        self.doc = Document()
        n = self.doc.styles['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(10.5)
        n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.15
        for i, sz in [(1, 17), (2, 13.5), (3, 12), (4, 11)]:
            s = self.doc.styles[f'Heading {i}']; s.font.name = 'Calibri'; s.font.size = Pt(sz)
            s.font.color.rgb = NAVY; s.font.bold = True

    def _bg(self, cell, hexc):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexc); tcPr.append(shd)

    def h(self, l, t): return self.doc.add_heading(t, level=l)

    def p(self, t, size=None, color=None, bold=False, italic=False, align=None):
        par = self.doc.add_paragraph(); r = par.add_run(t); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if align: par.alignment = align
        return par

    def b(self, t, lvl=0):
        par = self.doc.add_paragraph(t, style='List Bullet')
        par.paragraph_format.left_indent = Inches(0.3 + 0.25 * lvl); return par

    def num(self, t): return self.doc.add_paragraph(t, style='List Number')

    def pb(self): self.doc.add_page_break()

    def img(self, fn, w=6.6, cap=None):
        path = os.path.join(IMG, fn)
        if os.path.exists(path):
            self.doc.add_picture(path, width=Inches(w))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap:
                c = self.doc.add_paragraph(); r = c.add_run(cap); r.italic = True
                r.font.size = Pt(9); r.font.color.rgb = GREYC; c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, headers, rows, font=9, bgc='0b3d7b'):
        t = self.doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Table Grid'
        for i, hd in enumerate(headers):
            self._bg(t.rows[0].cells[i], bgc); r = t.rows[0].cells[i].paragraphs[0].add_run(hd)
            r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(font)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(font)
        return t

    def toc(self):
        par = self.doc.add_paragraph(); fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
        r = OxmlElement('w:r'); t = OxmlElement('w:t')
        t.text = 'Right-click and "Update Field" to build the Table of Contents.'
        r.append(t); fld.append(r); par._p.append(fld)

    def cover(self, title, subtitle):
        self.doc.add_paragraph('\n')
        self.p('GOVERNMENT OF INDIA', size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
        self.p('Ministry of Health & Family Welfare \u00b7 Directorate General of Health Services',
               size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREYC)
        self.p('Central Drugs Standard Control Organization (CDSCO)', size=13, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
        self.doc.add_paragraph('\n\n')
        self.p('DIGITAL DRUGS REGULATORY SYSTEM (DDRS)', size=23, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
        self.p(title, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=SAFFRON)
        self.p(subtitle, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREYC)
        self.doc.add_paragraph('\n\n\n')
        self.p('Document Version 1.0', size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.p('Date: ' + TODAY, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREYC)
        self.pb()
        self.h(1, 'Table of Contents'); self.toc(); self.pb()

    def save(self, fname):
        out = os.path.join(HERE, fname); self.doc.save(out)
        print('SAVED', fname, '| paras', len(self.doc.paragraphs), 'tables', len(self.doc.tables))
        return out


# =================================================================
# DOC 1 — DATA MIGRATION
# =================================================================
def build_migration():
    d = DDoc()
    d.cover('Data Migration Plan', 'Migrating Existing CDSCO Portals to the DDRS Portal')

    d.h(1, '1. Introduction')
    d.h(2, '1.1 Purpose')
    d.p('This document defines the strategy, methodology, mapping, governance and '
        'acceptance approach for migrating data from the existing CDSCO portals into the '
        'new Digital Drugs Regulatory System (DDRS). It ensures accuracy, completeness '
        'and integrity of migrated data while minimising downtime and disruption to '
        'ongoing regulatory activities.')
    d.h(2, '1.2 Scope')
    for t in ['Identification, extraction, cleansing, transformation, loading, validation '
              'and reconciliation of legacy data.',
              'Source systems: SUGAM, MD Online, SUGAM LABS, ONDLS and State portals.',
              'Master data (entities, products, technical persons, labs), transactional '
              'data (applications, licences, certificates), and documents.',
              'Incremental, phased migration aligned to the DDRS rollout.']:
        d.b(t)
    d.h(2, '1.3 Objectives & Principles')
    for t in ['Zero data loss and full traceability (every source record accounted for).',
              'Cleansed, de-duplicated, standardised data in DDRS.',
              'Reversible (rollback) and reconciled at each wave.',
              'Minimal downtime via parallel run and phased cutover.',
              'Security and confidentiality of sensitive data throughout.']:
        d.b(t)
    d.pb()

    d.h(1, '2. Source Systems Inventory')
    d.p('The legacy estate comprises the following CDSCO and State systems. Exact '
        'volumes and schemas are confirmed during the assessment phase with each system owner.')
    d.table(['Source System', 'Primary Content', 'Typical Formats', 'Target DDRS Area'],
            [['SUGAM (cdscoonline.gov.in)', 'Central licensing/permissions, registrations, CT', 'RDBMS, PDF, scans', 'applications, licensing, clinical-trials'],
             ['MD Online (cdscomdonline.gov.in)', 'Medical device & IVD manufacturing/import licences', 'RDBMS, PDF', 'applications, licensing, products'],
             ['SUGAM LABS (sugamlabs.gov.in)', 'Central lab LIMS: samples, tests, reports', 'RDBMS, instrument files', 'laboratory (samples, reports, BRC)'],
             ['ONDLS (statedrugs.gov.in)', 'State sales/manufacturing licences, COPP/GMP/MSC', 'RDBMS, PDF', 'applications, licensing, registry'],
             ['State portals (various)', 'State-specific licences & registrations', 'RDBMS/CSV/XLS', 'applications, registry, licensing'],
             ['Documents & scans', 'Supporting documents across portals', 'PDF/JPEG/TIFF', 'documents (with checksum/version)']])
    d.p('Reference data volume: existing portal utilisation is approximately 24.3 TB '
        '(as of March 2026), dominated by uploaded documents and scanned artefacts.')
    d.pb()

    d.h(1, '3. Migration Strategy')
    d.h(2, '3.1 Approach')
    d.p('Migration is incremental and wave-based, synchronised with the DDRS state-rollout '
        'plan. Each wave migrates a coherent slice (a portal/State and its dependent data) '
        'into a staging area, validates and reconciles it, then promotes it to production. '
        'A parallel-run window allows verification against the legacy system before cutover.')
    d.img('migration_flow.png', 6.9, 'Figure \u2014 DDRS data-migration pipeline.')
    d.h(2, '3.2 Migration Types')
    d.table(['Type', 'Use'],
            [['Big-bang (per slice)', 'Small, self-contained reference datasets'],
             ['Phased / incremental', 'Primary approach \u2014 by portal/State and entity group'],
             ['Parallel run', 'Critical transactional data verified against legacy before cutover'],
             ['Delta / catch-up', 'Capture records created in legacy during the migration window']])
    d.pb()

    d.h(1, '4. Migration Methodology')
    for i, (stage, desc) in enumerate([
        ('Assessment & Profiling', 'Inventory sources, profile data quality, define scope, volumes and risks per system.'),
        ('Extraction', 'Extract from source RDBMS/exports (API/DB/CSV); land in a secure staging area.'),
        ('Cleansing & De-duplication', 'Standardise formats, fix anomalies, de-duplicate entities/products, resolve orphan references.'),
        ('Transformation & Mapping', 'Map source schemas to DDRS entities/fields; derive reference numbers; normalise enums.'),
        ('Loading', 'Bulk-load into DDRS via repositories/ETL with referential integrity and batching.'),
        ('Validation', 'Row counts, checksums, business-rule and referential checks; sample audits.'),
        ('Reconciliation', 'Reconcile source vs target per entity; sign-off; record exceptions.'),
        ('Cutover & Hypercare', 'Switch users to DDRS; monitor; apply deltas; support.')], 1):
        d.h(3, '4.%d %s' % (i, stage)); d.p(desc)
    d.pb()

    d.h(1, '5. Data Mapping (Source \u2192 DDRS)')
    d.h(2, '5.1 Entity-Level Mapping')
    d.table(['Legacy Concept', 'Source', 'DDRS Entity', 'Notes'],
            [['Firm / Company', 'SUGAM/ONDLS/MD', 'organizations', 'De-dup by PAN/GSTIN/CIN'],
             ['Product / Brand', 'SUGAM/MD', 'products', 'Brand-name de-duplication applied'],
             ['Technical/competent person', 'ONDLS/SUGAM', 'technical_persons', 'Enforce one-entity engagement'],
             ['Testing laboratory', 'SUGAM LABS/ONDLS', 'laboratories', 'Central/State/Private typing'],
             ['Application / permission', 'SUGAM/MD/ONDLS', 'applications (+events)', 'Map status to DDRS state machine'],
             ['Licence / permission grant', 'SUGAM/MD/ONDLS', 'licenses', 'Preserve number; regenerate QR'],
             ['Certificate / NOC (COPP/FSC/MSC/NCC)', 'SUGAM/ONDLS', 'certificates', 'Map certType; set validity'],
             ['Sample / test report', 'SUGAM LABS', 'samples, test_reports', 'Map result to SQ/NSQ/Spurious'],
             ['Clinical trial', 'SUGAM', 'clinical_trials (+sites)', 'Link CTRI / EC approvals'],
             ['Payment / challan', 'SUGAM/ONDLS', 'payments', 'Status PAID; map gateway'],
             ['Uploaded document', 'All', 'documents', 'Store ref + checksum + version']])
    d.h(2, '5.2 Sample Field Mapping \u2014 Application')
    d.table(['Source Field (legacy)', 'DDRS Field', 'Transformation'],
            [['file_no / app_no', 'referenceNo', 'Preserve or re-key to CDSCO/<TYPE>/<YEAR>/<SEQ>'],
             ['applicant_id', 'organizationId', 'Resolve to migrated organization UUID'],
             ['form_type', 'type', 'Map to ApplicationType enum'],
             ['status_code', 'status', 'Map to ApplicationStatus (DRAFT\u2026ISSUED)'],
             ['submitted_on', 'submittedAt', 'Parse to ISO datetime'],
             ['fee_amount', 'feeAmount', 'Validate against fee rules'],
             ['product_category', 'productCategory', 'Map to ProductCategory enum'],
             ['state', 'stateCode', 'Map to 2-letter code']])
    d.pb()

    d.h(1, '6. Data Quality, Validation & Reconciliation')
    d.h(2, '6.1 Data Quality Dimensions')
    d.table(['Dimension', 'Check'],
            [['Completeness', 'Mandatory fields present; no truncation'],
             ['Accuracy', 'Values match source / reference data'],
             ['Consistency', 'Enums, codes and references normalised'],
             ['Uniqueness', 'No duplicates (entities, products, persons)'],
             ['Integrity', 'All foreign keys resolve; no orphans'],
             ['Validity', 'Conforms to DDRS validation rules']])
    d.h(2, '6.2 Reconciliation')
    for t in ['Per-entity source vs target row counts with variance log.',
              'Checksum/hash comparison for documents.',
              'Business-rule validation (e.g., licence validity, status legality).',
              'Statistical and sample-based manual audits with sign-off.']:
        d.b(t)
    d.h(2, '6.3 Acceptance Criteria')
    for t in ['100% of in-scope records migrated or formally exception-listed.',
              'Reconciliation variance within agreed tolerance (target 0 for critical data).',
              'No unresolved referential-integrity errors.',
              'CDSCO sign-off per wave before production promotion.']:
        d.b(t)
    d.pb()

    d.h(1, '7. Cutover, Rollback & Tools')
    d.h(2, '7.1 Cutover & Parallel Run')
    d.p('Each wave runs DDRS in parallel with the legacy system for a defined window; '
        'delta records are captured and applied; users are switched at cutover with hypercare support.')
    d.h(2, '7.2 Rollback')
    d.p('Every load is reversible: staging snapshots and pre-load backups allow a wave to '
        'be rolled back without affecting previously migrated, signed-off data.')
    d.h(2, '7.3 Tooling')
    d.table(['Concern', 'Approach'],
            [['Extraction', 'Source DB connectors / API exports / CSV-TSV'],
             ['Staging', 'Isolated schema with data-quality rules and audit'],
             ['Transform/Load', 'Scripted ETL into DDRS (TypeORM/SQL), batched'],
             ['Validation', 'Automated reconciliation reports + dashboards'],
             ['Documents', 'Object storage with checksum + version capture']])
    d.pb()

    d.h(1, '8. Governance, RACI, Risks & Schedule')
    d.h(2, '8.1 RACI')
    d.table(['Activity', 'CDSCO', 'SSP', 'Source Owner', 'DBA'],
            [['Source access & schemas', 'A', 'C', 'R', 'C'],
             ['Profiling & mapping', 'A', 'R', 'C', 'C'],
             ['Cleansing rules', 'A', 'R', 'C', 'I'],
             ['ETL build & load', 'I', 'R', 'I', 'C'],
             ['Validation & reconciliation', 'A', 'R', 'C', 'C'],
             ['Cutover sign-off', 'A', 'C', 'C', 'I']])
    d.h(2, '8.2 Risks & Mitigations')
    d.table(['Risk', 'Mitigation'],
            [['Poor legacy data quality', 'Early profiling; cleansing rules; exception handling'],
             ['Schema ambiguity', 'Workshops with system owners; iterative mapping'],
             ['Large document volume (24.3 TB)', 'Phased document migration; checksum verification'],
             ['Downtime risk', 'Parallel run + delta capture + scheduled cutover windows'],
             ['Referential gaps', 'Dependency-ordered loading; orphan resolution'],
             ['Sensitive data exposure', 'Encryption, access control, masked non-prod data']])
    d.h(2, '8.3 Indicative Schedule')
    d.table(['Wave', 'Scope', 'Window'],
            [['W0', 'Assessment, profiling, mapping, tooling', 'Months 1\u20132'],
             ['W1', 'Central master + SUGAM/MD core', 'Months 2\u20134'],
             ['W2', 'SUGAM LABS + high-volume States (ONDLS)', 'Months 4\u20137'],
             ['W3', 'Remaining States/UTs + documents', 'Months 7\u201310'],
             ['W4', 'Delta, reconciliation close-out, archive', 'Months 10\u201312']])
    d.pb()
    d.h(1, 'Appendix A \u2014 Reconciliation Report Template')
    d.table(['Entity', 'Source count', 'Target count', 'Variance', 'Exceptions', 'Status'],
            [['organizations', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014'],
             ['products', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014'],
             ['applications', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014'],
             ['licenses', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014'],
             ['samples', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014']])
    d.h(1, 'Appendix B \u2014 Mapping Specification Template')
    d.table(['Source field', 'Type', 'DDRS field', 'Type', 'Rule', 'Default'],
            [['\u2014'] * 6])
    d.p('\n'); d.p('\u2014 End of Document \u2014', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=NAVY)
    return d.save('DDRS_Data_Migration_Plan.docx')


# =================================================================
# DOC 2 — TRANSITION / REFACTORING / MIGRATION / UPGRADATION / MAINTENANCE
# =================================================================
def build_transition():
    d = DDoc()
    d.cover('Transition, Refactoring, Migration, Upgradation & Maintenance Plan',
            'Operating and Evolving the DDRS Portal Post Go-Live')
    d.h(1, '1. Introduction & Objectives')
    d.p('This plan describes how the DDRS portal will be transitioned, refactored, '
        'migrated, upgraded and maintained after Go-Live, ensuring continuous, secure and '
        'high-quality service over the operations period and a clean eventual exit/handover.')
    for t in ['Sustain agreed service levels (availability, performance, security).',
              'Continuously improve the platform without disrupting operations.',
              'Keep the codebase clean, modular and microservices-ready.',
              'Maintain complete, current documentation and transferable knowledge.']:
        d.b(t)
    d.img('transition_timeline.png', 6.9, 'Figure \u2014 Transition & post-Go-Live lifecycle.')
    d.pb()

    d.h(1, '2. Transition & Knowledge Acquisition')
    for t in ['Structured knowledge acquisition: architecture, code, data, runbooks.',
              'Shadowing and reverse-KT with outgoing teams (where applicable).',
              'Environment, repository, CI/CD and credentials handover with inventory.',
              'Baseline of SLAs, backlog, known issues and technical debt.']:
        d.b(t)
    d.pb()

    d.h(1, '3. Refactoring Strategy')
    d.p('DDRS is a modular monolith that is microservices-ready: each domain is isolated '
        'behind a documented OpenAPI contract. Refactoring is continuous and contract-safe.')
    d.table(['Theme', 'Approach'],
            [['Tech-debt management', 'Debt register; allocate capacity each iteration'],
             ['Modular boundaries', 'Keep domains decoupled; extract to services when load demands'],
             ['Code quality', 'Typed code, lint gates, tests, peer review'],
             ['API versioning', 'Backward-compatible changes; deprecation policy'],
             ['Performance', 'Profiling, indexing, query/pagination tuning'],
             ['Reusability', 'Shared utilities (pagination, reference/QR, audit)']])
    d.pb()

    d.h(1, '4. Migration (Environment & Data)')
    for t in ['Development SQLite \u2192 production PostgreSQL via DB_TYPE config (portable schema).',
              'Schema evolution via controlled migrations; no destructive changes in prod.',
              'Data migration/backfills run in maintenance windows with backups and rollback.',
              'Zero/low-downtime techniques (expand-contract schema changes).']:
        d.b(t)
    d.pb()

    d.h(1, '5. Upgradation & Release Management')
    d.table(['Aspect', 'Practice'],
            [['Source control', 'Trunk-based with protected main; PR review'],
             ['CI/CD', 'Lint \u2192 test \u2192 build \u2192 image \u2192 deploy pipeline'],
             ['Deployment', 'Blue-green / canary; health checks; instant rollback'],
             ['Feature flags', 'Decouple deploy from release; controlled exposure'],
             ['Versioning', 'Semantic versioning; documented release notes'],
             ['Change windows', 'Scheduled, communicated; emergency-change path']])
    d.pb()

    d.h(1, '6. Maintenance Services')
    d.table(['Type', 'Definition', 'Examples'],
            [['Corrective', 'Fix defects', 'Bug fixes, hotfixes'],
             ['Adaptive', 'Adapt to change', 'Rule/Act updates, new integrations, OS/lib updates'],
             ['Perfective', 'Improve', 'Performance, UX, new reports'],
             ['Preventive', 'Prevent issues', 'Refactoring, patching, capacity tuning']])
    d.h(2, '6.1 Support Model')
    d.table(['Tier', 'Scope'],
            [['L1 Helpdesk', 'IVRS/ticket intake, triage, FAQs, password/access'],
             ['L2 Application', 'Functional issues, configuration, data fixes'],
             ['L3 Engineering', 'Code defects, performance, integrations, DB'],
             ['Vendor/3rd-party', 'Cloud, gateways, external systems']])
    d.h(2, '6.2 Indicative SLAs')
    d.table(['Severity', 'Response', 'Resolution'],
            [['S1 Critical', '15 min', '4 h'], ['S2 High', '30 min', '8 h'],
             ['S3 Medium', '4 h', '3 business days'], ['S4 Low', '1 business day', '10 business days']])
    d.pb()

    d.h(1, '7. Operations: Monitoring, Security, Backup & DR')
    d.h(2, '7.1 Monitoring & Observability')
    for t in ['Health endpoint probes; uptime and latency monitoring.',
              'Structured logs, metrics and tracing; alerting on thresholds.',
              'Audit-trail and integration-log review for anomalies.']:
        d.b(t)
    d.h(2, '7.2 Security Operations')
    for t in ['Regular CERT-In empanelled audits; vulnerability scanning and patching.',
              'Secret rotation (JWT, DB, gateway keys); least-privilege access.',
              'Incident response plan with defined severities and comms.']:
        d.b(t)
    d.h(2, '7.3 Backup, DR & BCP')
    d.table(['Control', 'Target'],
            [['Backup frequency', 'Daily full + frequent incrementals (PostgreSQL)'],
             ['RPO', '\u2264 15 minutes (PITR)'], ['RTO', '\u2264 1 hour'],
             ['DR', 'Cross-zone/region restore drills; runbook-tested'],
             ['Retention', 'Per policy; audit logs archived off-site']])
    d.pb()

    d.h(1, '8. Change Management, Documentation & Governance')
    for t in ['Change Control Board reviews significant changes (CCB).',
              'Living documentation: SAD, runbooks, API docs, change logs kept current.',
              'Periodic reviews (Tech Ops, Architecture Review) and reporting.']:
        d.b(t)
    d.h(2, '8.1 Governance RACI')
    d.table(['Activity', 'CDSCO', 'SSP', 'CCB'],
            [['Roadmap & priorities', 'A', 'C', 'C'], ['Releases', 'A', 'R', 'C'],
             ['Emergency changes', 'A', 'R', 'I'], ['SLA reporting', 'A', 'R', 'I']])
    d.pb()

    d.h(1, '9. Exit / Transition Management')
    d.p('On contract end or transition to a new provider, the SSP ensures a clean, '
        'complete handover with no service disruption.')
    for t in ['Updated transition plan and complete documentation handover.',
              'Full source code, configuration, infrastructure-as-code and data export.',
              'Inventory of assets, credentials and access; knowledge-transfer sessions.',
              'Walk-throughs/demos; defined acceptance of handover by CDSCO/new vendor.']:
        d.b(t)
    d.h(1, '10. Risks & KPIs')
    d.table(['Risk', 'Mitigation'],
            [['Knowledge loss', 'Documentation + cross-training + recorded KT'],
             ['Regression on change', 'Automated tests + canary + rollback'],
             ['Security drift', 'Continuous scanning + patch cadence + audits'],
             ['Scope creep', 'CCB-governed change control']])
    d.table(['KPI', 'Target'],
            [['Availability', '\u2265 99.5% monthly'], ['S1 resolution within SLA', '\u2265 95%'],
             ['Change success rate', '\u2265 98%'], ['Mean time to recover', '\u2264 1 h'],
             ['Documentation currency', '100% per release']])
    d.p('\n'); d.p('\u2014 End of Document \u2014', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=NAVY)
    return d.save('DDRS_Transition_Maintenance_Plan.docx')


# =================================================================
# DOC 3 — ADVANCED ANALYTICS & INTELLIGENT AUTOMATION — PROJECT CITATION FORMAT
# =================================================================
def build_analytics_citation():
    d = DDoc()
    d.cover('Advanced Analytics & Intelligent Automation',
            'Project Citation Format & Reference Submission Template')
    d.h(1, '1. Purpose & Instructions')
    d.p('This format is used to cite a prior or current project demonstrating capability '
        'in advanced analytics and intelligent automation. Complete one citation per '
        'project. Provide verifiable details and attach supporting evidence '
        '(work order, completion/performance certificate or client reference). Incomplete '
        'or unverifiable citations may not be considered.')
    for t in ['Use one form per cited project.',
              'State quantifiable outcomes (volumes, % improvement, time saved).',
              'Clearly identify the analytics/automation techniques applied.',
              'Attach client-issued evidence for each citation.']:
        d.b(t)
    d.pb()

    d.h(1, '2. Project Citation Form')
    d.table(['#', 'Field', 'Response'],
            [['1', 'Project title', ''], ['2', 'Client / Organisation', ''],
             ['3', 'Sector / Domain', ''], ['4', 'Client contact (name, designation, email/phone)', ''],
             ['5', 'Order/Contract value (INR)', ''], ['6', 'Start date', ''], ['7', 'End date / Status', ''],
             ['8', 'Role of the bidder (prime/sub)', ''], ['9', 'Team size & key roles', ''],
             ['10', 'Brief scope (3\u20135 lines)', ''],
             ['11', 'Advanced analytics techniques used', ''],
             ['12', 'Intelligent automation techniques used', ''],
             ['13', 'Data scale (volume/velocity/variety)', ''],
             ['14', 'Tools / platforms / models', ''],
             ['15', 'Integration with enterprise systems', ''],
             ['16', 'Quantifiable outcomes / benefits', ''],
             ['17', 'Governance, security & privacy measures', ''],
             ['18', 'Evidence attached (type & reference)', '']], font=9)
    d.pb()

    d.h(1, '3. Qualifying Criteria & Scoring Guidance')
    d.table(['Criterion', 'What evidences strength', 'Weight'],
            [['Relevance to regulatory/Govt domain', 'Similar domain, scale, stakeholders', 'High'],
             ['Advanced analytics depth', 'Predictive/prescriptive, dashboards, benchmarking indices', 'High'],
             ['Intelligent automation', 'Workflow automation, NLP/RAG, anomaly detection, agents', 'High'],
             ['Outcomes', 'Quantified efficiency/quality/compliance gains', 'High'],
             ['Data scale & integration', 'Large/heterogeneous data; enterprise integration', 'Medium'],
             ['Governance', 'Security, privacy (DPDP), model governance, auditability', 'Medium']])
    d.h(2, '3.1 Techniques Checklist')
    for t in ['Descriptive & diagnostic analytics (dashboards, MIS, drill-down).',
              'Predictive analytics (risk scoring, forecasting).',
              'Prescriptive analytics / optimisation.',
              'Benchmarking indices and composite scoring.',
              'NLP / RAG / document intelligence.',
              'Anomaly / fraud detection.',
              'Workflow / robotic / agentic automation.',
              'MLOps and model governance.']:
        d.b(t)
    d.pb()

    d.h(1, '4. Illustrative Worked Example (DDRS)')
    d.p('The DDRS programme itself exemplifies advanced analytics and intelligent '
        'automation in a national regulatory context; the filled citation below is '
        'illustrative of the expected level of detail.')
    d.table(['Field', 'Illustrative Response (DDRS)'],
            [['Project title', 'Digital Drugs Regulatory System (DDRS) \u2014 analytics & automation'],
             ['Client', 'CDSCO, Ministry of Health & Family Welfare, Government of India'],
             ['Sector', 'Drug regulation / e-Governance / DPI'],
             ['Scope', 'Unified regulatory platform with MIS dashboards, SHRESTH state-benchmarking index, custom report builder, and intelligent automation across the application lifecycle.'],
             ['Advanced analytics', 'Role-based dashboards; trend analysis; SHRESTH composite index across 36 States/UTs; enforcement/quality surveillance analytics; custom report builder with export.'],
             ['Intelligent automation', 'Generic workflow engine; random/auto masked work-allocation; Timeline Review System (SLA ageing & alerts); fee auto-calculation; auto-issuance of licences/certificates with QR; automated notifications/alerts.'],
             ['Data scale', '60,000+ seeded demonstration records across 37 entities; designed for ~10 lakh stakeholders and ~1,000 concurrent users.'],
             ['Tools / platforms', 'NestJS, TypeScript, TypeORM, PostgreSQL, React, Recharts; AI-ready APIs for RAG/agents.'],
             ['Integration', 'Hub of 35 government/external systems (Aadhaar, GST, Customs, Bharat Kosh, CTRI, NPPA, ABDM, IPC, NIB \u2026).'],
             ['Outcomes (illustrative)', 'Paperless, time-bound approvals; SLA visibility; surveillance-driven enforcement; transparent public verification.'],
             ['Governance', 'RBAC (19 roles), immutable audit trail, DPDP-aligned access, ISO 27001 alignment, CERT-In audit.']])
    d.pb()

    d.h(1, '5. Supporting-Evidence Checklist')
    for t in ['Work order / purchase order / contract.',
              'Completion or performance certificate from the client.',
              'Client reference letter or contactable reference.',
              'Architecture/outcome summary or case study (if permitted).']:
        d.b(t)
    d.h(1, '6. Declaration')
    d.p('We certify that the information provided in each citation is true, accurate and '
        'verifiable, and that supporting evidence is enclosed.')
    d.table(['Field', 'Value'],
            [['Authorised signatory', ''], ['Designation', ''], ['Organisation', ''],
             ['Date', ''], ['Signature / Seal', '']])
    d.p('\n'); d.p('\u2014 End of Document \u2014', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=NAVY)
    return d.save('DDRS_Advanced_Analytics_Project_Citation_Format.docx')


# =================================================================
# DOC 4 — INTELLIGENT SYSTEMS ARCHITECTURE & REGULATORY USE CASES
# =================================================================
def build_intelligent_arch():
    d = DDoc()
    d.cover('Intelligent Systems Architecture & Regulatory Use Cases',
            'AI-Ready Design and Applied AI for Drug Regulation')
    d.h(1, '1. Introduction')
    d.p('AI-enabled capabilities are envisaged to enhance the efficiency, accuracy and '
        'responsiveness of DDRS. Per the programme\u2019s provisional AI scope, AI models and '
        'agentic capabilities are developed separately; DDRS\u2019s responsibility is to be '
        'architected as AI-ready so externally developed AI can be integrated without '
        're-engineering the core. This document describes that intelligent-systems '
        'architecture and the regulatory AI use cases it enables.')
    for t in ['All platform services exposed via standardised, versioned, documented APIs.',
              'Architecture supports plug-in integration of external AI microservices and agents.',
              'Infrastructure accounts for future AI inference and orchestration workloads.',
              'Data models, schemas and workflows documented for AI training and agent onboarding.']:
        d.b(t)
    d.pb()

    d.h(1, '2. AI-Ready Reference Architecture')
    d.img('ai_architecture.png', 6.9, 'Figure \u2014 AI-ready architecture (plug-in intelligence).')
    d.table(['Component', 'Role'],
            [['DDRS OpenAPI surface', 'Single, versioned contract over all modules and data'],
             ['AI Gateway / Orchestration', 'Routes requests to AI services; hosts agents (Model Context Protocol)'],
             ['AI microservices', 'RAG assistant, summarisation, classification, comparison, etc.'],
             ['Vector store / knowledge index', 'Embeddings of guidelines, dossiers, SOPs for retrieval'],
             ['Model registry & MLOps', 'Versioning, deployment, monitoring, evaluation of models'],
             ['Governance & guardrails', 'PII protection, human-in-loop, audit, policy enforcement']])
    d.h(2, '2.1 Integration Principles')
    for t in ['Loose coupling: AI services consume DDRS APIs; no direct DB coupling.',
              'Asynchronous, event-driven where suitable; idempotent operations.',
              'Standard protocols (REST/JSON, Model Context Protocol for agents).',
              'Versioned contracts and feature flags for safe rollout.',
              'Every AI action audited; human-in-the-loop for regulatory decisions.']:
        d.b(t)
    d.pb()

    d.h(1, '3. Regulatory AI Use-Case Map')
    d.img('ai_usecases.png', 6.6, 'Figure \u2014 Regulatory AI use-case map.')
    d.h(2, '3.1 Use-Case Catalogue')
    UC = [
     ('RAG conversational assistant', 'Answer applicant/officer queries grounded in CDSCO guidelines, Acts & SOPs', 'Query + knowledge index', 'Cited answer', 'RAG over vector store', 'Faster, consistent guidance'),
     ('Document summarisation', 'Summarise dossiers, reports, PSURs', 'Document', 'Structured summary', 'LLM summarisation', 'Reviewer time saved'),
     ('Completeness assessment', 'Check application/dossier completeness vs checklist', 'Application + rules', 'Gap report', 'Classification + rules', 'Fewer query cycles'),
     ('Classification & routing', 'Auto-classify applications/complaints and route', 'Application/complaint', 'Category + route', 'Text classification', 'Faster allocation'),
     ('Document comparison', 'Compare versions/labels/specifications', 'Two documents', 'Diff & risk flags', 'Semantic diff', 'Detect discrepancies'),
     ('Inspection report generation', 'Draft inspection reports from findings', 'Findings + template', 'Draft report', 'LLM generation', 'Inspector time saved'),
     ('AI anonymisation', 'Redact PII for sharing/analytics', 'Records/documents', 'Anonymised data', 'PII detection', 'DPDP-safe sharing'),
     ('Fraud / anomaly analytics', 'Detect duplicate brands, suspicious patterns, NSQ risk', 'Transactions/registries', 'Risk scores/alerts', 'Anomaly detection', 'Proactive surveillance'),
     ('Duplicate detection', 'Detect duplicate licences/brand names/batches', 'Registry data', 'Duplicate clusters', 'Similarity/embeddings', 'Data integrity'),
     ('Risk-based inspection scoring', 'Prioritise inspections by risk', 'History + signals', 'Risk ranking', 'Predictive model', 'Efficient enforcement'),
     ('Multi-agent workflows (MCP)', 'Orchestrate multi-step regulatory tasks', 'Task + tools/APIs', 'Completed actions', 'Agentic orchestration', 'Automation at scale'),
     ('Alert generation', 'Draft drug/cosmetic alerts from signals', 'Enforcement/lab data', 'Draft alert', 'LLM + rules', 'Timely public safety'),
    ]
    d.table(['Use case', 'Description', 'Inputs', 'Outputs', 'Technique', 'Benefit'],
            [list(u) for u in UC], font=8)
    d.pb()

    d.h(1, '4. Data Readiness & MLOps')
    for t in ['Documented schemas (37 entities) and workflows enable supervised learning and retrieval.',
              'Document & knowledge stores feed the vector index for RAG.',
              'MLOps: model registry, versioned deployment, monitoring, drift detection, evaluation.',
              'Feedback loops capture human corrections to improve models.']:
        d.b(t)
    d.h(1, '5. Governance, Ethics & Guardrails')
    d.table(['Concern', 'Control'],
            [['Accuracy / hallucination', 'Retrieval grounding + citations + confidence thresholds'],
             ['Human oversight', 'Human-in-the-loop for all regulatory decisions; AI assists, not decides'],
             ['Privacy (DPDP)', 'PII minimisation/anonymisation; access control; purpose limitation'],
             ['Auditability', 'All AI invocations logged with inputs/outputs/version'],
             ['Bias & fairness', 'Evaluation datasets; monitoring; review'],
             ['Security', 'Prompt-injection defences; sandboxed tools; least privilege']])
    d.h(1, '6. Roadmap & Maturity')
    d.table(['Stage', 'Capability'],
            [['Foundation (now)', 'AI-ready APIs, documented data, dashboards & rule-based automation'],
             ['Assistive', 'RAG assistant, summarisation, completeness checks'],
             ['Augmented', 'Classification, anomaly/fraud analytics, risk-based inspection'],
             ['Autonomous (governed)', 'Multi-agent workflows via MCP with human-in-loop']])
    d.h(1, '7. Risks & Mitigations')
    d.table(['Risk', 'Mitigation'],
            [['Over-reliance on AI', 'Human-in-loop; AI as decision support only'],
             ['Data quality', 'Validation, cleansing, monitoring'],
             ['Model drift', 'Continuous evaluation and retraining'],
             ['Privacy exposure', 'Anonymisation, access control, DPDP compliance'],
             ['Integration risk', 'Versioned APIs, contract tests, feature flags']])
    d.h(1, 'Appendix A \u2014 Sample Agent Workflow (MCP)')
    for t in ['User asks: "Is application CDSCO/ML/2026/000123 ready for approval?"',
              'Agent retrieves the application + checklist via DDRS APIs.',
              'Completeness service flags missing documents; RAG cites the relevant rule.',
              'Agent drafts a query letter (human reviews and sends).',
              'All steps and tool calls are logged to the audit trail.']:
        d.num(t)
    d.p('\n'); d.p('\u2014 End of Document \u2014', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=NAVY)
    return d.save('DDRS_Intelligent_Systems_Architecture_and_Use_Cases.docx')


if __name__ == '__main__':
    build_migration()
    build_transition()
    build_analytics_citation()
    build_intelligent_arch()
    print('ALL 4 DOCS BUILT')
