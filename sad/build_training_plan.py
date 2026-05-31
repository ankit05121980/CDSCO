#!/usr/bin/env python3
"""Build the DDRS Capacity Building & Training Plan (.docx)."""
import os, datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(__file__)
IMG = os.path.join(HERE, 'images')

NAVY = RGBColor(0x0b, 0x3d, 0x7b); SAFFRON = RGBColor(0xc0, 0x60, 0x10)
GREEN = RGBColor(0x0e, 0x66, 0x06); GREYC = RGBColor(0x55, 0x60, 0x6b)

doc = Document()
n = doc.styles['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(10.5)
n.paragraph_format.space_after = Pt(6); n.paragraph_format.line_spacing = 1.15
for i, sz in [(1, 18), (2, 14), (3, 12), (4, 11)]:
    s = doc.styles[f'Heading {i}']; s.font.name = 'Calibri'; s.font.size = Pt(sz)
    s.font.color.rgb = NAVY; s.font.bold = True


def bg(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexc); tcPr.append(shd)


def h(l, t): return doc.add_heading(t, level=l)


def para(t, size=None, color=None, bold=False, italic=False, align=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    return p


def bullet(t, lvl=0):
    p = doc.add_paragraph(t, style='List Bullet'); p.paragraph_format.left_indent = Inches(0.3 + 0.25 * lvl); return p


def num(t): return doc.add_paragraph(t, style='List Number')


def pb(): doc.add_page_break()


def img(fn, w=6.6, cap=None):
    p = os.path.join(IMG, fn)
    if os.path.exists(p):
        doc.add_picture(p, width=Inches(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if cap:
            c = doc.add_paragraph(); r = c.add_run(cap); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREYC
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(headers, rows, font=9, bgc='0b3d7b'):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, hd in enumerate(headers):
        bg(t.rows[0].cells[i], bgc); r = t.rows[0].cells[i].paragraphs[0].add_run(hd)
        r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(font)
    return t


def toc():
    p = doc.add_paragraph(); fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'Right-click and "Update Field" to build the Table of Contents.'
    r.append(t); fld.append(r); p._p.append(fld)


# ---------------- COVER ----------------
doc.add_paragraph('\n')
para('GOVERNMENT OF INDIA', size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
para('Ministry of Health & Family Welfare · Directorate General of Health Services', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREYC)
para('Central Drugs Standard Control Organization (CDSCO)', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
doc.add_paragraph('\n\n')
para('DIGITAL DRUGS REGULATORY SYSTEM (DDRS)', size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)
para('CAPACITY BUILDING & TRAINING PLAN', size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=SAFFRON)
doc.add_paragraph('\n\n\n')
para('Document Version 1.0', size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Date: ' + datetime.date.today().strftime('%d %B %Y'), size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREYC)
pb()

h(1, 'Document Control')
table(['Version', 'Date', 'Author', 'Description'],
      [['0.5', '—', 'Training & Change Mgmt', 'Draft plan'],
       ['1.0', datetime.date.today().strftime('%d %b %Y'), 'Training & Change Mgmt', 'Baseline release']])
h(2, 'Purpose')
para('This Capacity Building & Training Plan defines the strategy, curricula, schedules, '
     'delivery methods, materials, assessment and governance for equipping all DDRS '
     'stakeholders — CDSCO officers, State Licensing Authorities, testing laboratories, '
     'industry and the public — with the knowledge and skills to use the platform '
     'effectively. It covers a phased rollout across all 36 States/UTs and the full '
     'three-year operations period.')
h(2, 'Scope')
for t in ['Role-based training for all 19 DDRS user roles across five portals.',
          'Train-the-trainer (ToT) and master-trainer cadre development.',
          'E-learning, instructor-led, hands-on lab, webinars and on-site sessions.',
          'Training materials: manuals, quick-reference guides, videos, simulations.',
          'State onboarding programme and phased rollout schedule.',
          'Assessment, certification, feedback and continuous improvement.',
          'Helpdesk, knowledge base and post-go-live support.']:
    bullet(t)
pb()
h(1, 'Table of Contents'); toc(); pb()
print('front matter done')

# ============ 1. STRATEGY ============
h(1, '1. Training Strategy & Approach')
para('The capacity-building programme follows a blended, role-based, competency-driven '
     'approach designed for a geographically distributed, multi-stakeholder audience. It '
     'uses the ADDIE instructional design model (Analyse, Design, Develop, Implement, '
     'Evaluate) and a train-the-trainer multiplier model to scale across 36 States/UTs.')
h(2, '1.1 Objectives')
for t in ['Enable every user role to perform its DDRS tasks confidently and correctly.',
          'Reduce support tickets and processing errors through proficiency.',
          'Build a self-sustaining master-trainer cadre within CDSCO and the States.',
          'Drive adoption, data quality and timely (SLA-compliant) processing.',
          'Institutionalise continuous learning and knowledge management.']:
    bullet(t)
h(2, '1.2 Guiding Principles')
for t in ['Role-based: curricula tailored to each role\u2019s portal and tasks.',
          'Blended: e-learning + instructor-led + hands-on sandbox practice.',
          'Hands-on first: practice on a seeded training environment, not slides alone.',
          'Multilingual: English and Hindi, with regional language support where needed.',
          'Accessible: GIGW-aligned materials; self-paced options for flexibility.',
          'Measured: assessments, certification and Kirkpatrick-level evaluation.']:
    bullet(t)
h(2, '1.3 Blended Learning Model')
table(['Mode', 'Use', 'Share'],
      [['E-learning (self-paced)', 'Foundation, refreshers, just-in-time micro-modules', '40%'],
       ['Instructor-led (virtual)', 'Role tracks, Q&A, scenario walkthroughs', '30%'],
       ['Hands-on lab (sandbox)', 'Guided practice on seeded environment', '20%'],
       ['Webinars / workshops', 'Updates, advanced topics, communities of practice', '10%']])
pb()

# ============ 2. STAKEHOLDER & ROLE ANALYSIS ============
h(1, '2. Stakeholder & Training-Needs Analysis')
para('Training needs are mapped to the 19 DDRS roles and the five portals. The matrix '
     'below shows the target proficiency per role across the principal modules '
     '(F=Foundation, I=Intermediate, A=Advanced).')
img('rbac.png', 6.4, 'Roles mapped to portals (drives the role-based curricula).')
table(['Role', 'Portal', 'Applications', 'Inspections/Enf.', 'Lab', 'Vigilance', 'Analytics', 'Admin'],
      [['DCGI / ADC', 'CDSCO', 'A', 'A', 'I', 'A', 'A', 'I'],
       ['Review Officer', 'CDSCO', 'A', 'I', 'F', 'I', 'I', '-'],
       ['Drug Inspector', 'CDSCO', 'I', 'A', 'I', 'F', 'F', '-'],
       ['Port Officer', 'CDSCO', 'A', 'F', '-', '-', 'F', '-'],
       ['State Licensing Authority', 'State', 'A', 'A', 'I', 'F', 'A', '-'],
       ['State Drug Inspector', 'State', 'I', 'A', 'I', 'F', 'F', '-'],
       ['Lab Manager', 'Lab', '-', 'F', 'A', 'F', 'F', '-'],
       ['Lab Analyst', 'Lab', '-', '-', 'A', '-', '-', '-'],
       ['Manufacturer/Importer', 'Industry', 'A', 'F', 'F', 'I', 'F', '-'],
       ['Wholesaler/Retailer', 'Industry', 'I', '-', '-', '-', '-', '-'],
       ['CRO / Ethics Committee', 'Industry', 'I', '-', '-', 'I', '-', '-'],
       ['Blood Centre', 'Industry', 'I', 'F', 'I', 'I', '-', '-'],
       ['Technical Person', 'Industry', 'F', '-', '-', '-', '-', '-'],
       ['Super Admin', 'Admin', 'I', 'I', 'I', 'I', 'A', 'A'],
       ['Public User', 'Public', 'F', '-', '-', '-', '-', '-']], font=8.5)
para('Approximate audience sizing: ~6,500 internal users (CDSCO + States + labs) and a '
     'large external base (~10 lakh stakeholders) trained primarily through self-paced '
     'e-learning, webinars and the public knowledge base.')
pb()

# ============ 3. CURRICULUM ============
h(1, '3. Curriculum & Course Catalogue')
para('The curriculum is organised into a common Foundation track, role-specific tracks '
     'mapped to portals, and specialised Administration, Integration and Security tracks.')
h(2, '3.1 Curriculum Structure')
img('portal_nav.png', 6.4, 'Module/portal map underpinning the role tracks.')
h(2, '3.2 Course Catalogue')
COURSES = [
 ('DDRS-F01', 'DDRS Overview & Navigation', 'All roles', '1.5 h', 'E-learning', 'None'),
 ('DDRS-F02', 'Login, MFA (OTP/Aadhaar) & Profile', 'All roles', '0.5 h', 'E-learning', 'None'),
 ('DDRS-F03', 'Dashboards, Search, Filters & Notifications', 'All roles', '1 h', 'E-learning', 'F01'),
 ('DDRS-F04', 'Data Protection, Audit & Responsible Use (DPDP)', 'All roles', '1 h', 'E-learning', 'F01'),
 ('DDRS-IND01', 'Creating & Submitting Applications', 'Industry', '2 h', 'ILT + Lab', 'F01-F03'),
 ('DDRS-IND02', 'Fee Payment (Bharat Kosh/Treasury) & Receipts', 'Industry', '1 h', 'Lab', 'IND01'),
 ('DDRS-IND03', 'Licences, Certificates, NOCs & QR Verification', 'Industry', '1.5 h', 'ILT', 'IND01'),
 ('DDRS-IND04', 'Renewals, Endorsements & Post-Approval Changes', 'Industry', '1.5 h', 'ILT + Lab', 'IND01'),
 ('DDRS-IND05', 'Clinical Trials & Site Management', 'CRO/Sponsor', '2 h', 'ILT', 'IND01'),
 ('DDRS-IND06', 'Vigilance Reporting (SAE/AEFI) & E2B/ICSR Import', 'Industry/PV', '2 h', 'ILT + Lab', 'F01'),
 ('DDRS-IND07', 'Returns Filing & Supply-Chain Batch/QR', 'Industry', '1.5 h', 'Lab', 'IND01'),
 ('DDRS-CEN01', 'Work Queue, Allocation & Timeline Review (TRS)', 'CDSCO/State', '2 h', 'ILT + Lab', 'F03'),
 ('DDRS-CEN02', 'Application Review, Queries & Decisions', 'CDSCO/State', '2.5 h', 'ILT + Lab', 'CEN01'),
 ('DDRS-CEN03', 'Issuing Licences, Certificates & NOCs', 'CDSCO/State', '1.5 h', 'Lab', 'CEN02'),
 ('DDRS-CEN04', 'Inspections: Scheduling, Forms & Joint Reports', 'Inspectors', '2.5 h', 'ILT + Lab', 'F03'),
 ('DDRS-CEN05', 'Enforcement: Sampling, NSQ/Spurious, Recalls, Court', 'CDSCO/State', '2.5 h', 'ILT', 'F03'),
 ('DDRS-CEN06', 'Vigilance Assessment & PSUR Review', 'CDSCO', '2 h', 'ILT', 'F03'),
 ('DDRS-LAB01', 'Sample Receipt & Workflow (LIMS)', 'Lab', '2 h', 'ILT + Lab', 'F03'),
 ('DDRS-LAB02', 'Testing, Reports & Result Capture', 'Lab Analyst', '2.5 h', 'Lab', 'LAB01'),
 ('DDRS-LAB03', 'Batch Release (BRC), SLP Scrutiny & Reference Standards', 'Lab Manager', '2 h', 'ILT', 'LAB01'),
 ('DDRS-ANA01', 'MIS Dashboards & Custom Reports (CSV)', 'CDSCO/State', '1.5 h', 'ILT', 'F03'),
 ('DDRS-ANA02', 'SHRESTH State Benchmarking Index', 'Authorities', '1 h', 'Webinar', 'ANA01'),
 ('DDRS-ADM01', 'User & Role Administration (RBAC)', 'Admin', '2 h', 'ILT + Lab', 'F04'),
 ('DDRS-ADM02', 'Entity/Lab Onboarding & Master Data', 'Admin', '1.5 h', 'Lab', 'ADM01'),
 ('DDRS-ADM03', 'Audit Trail & Investigations', 'Admin/CDSCO', '1.5 h', 'ILT', 'F04'),
 ('DDRS-INT01', 'Integrations Hub & External Systems (35)', 'Admin/CDSCO', '2 h', 'ILT', 'ADM01'),
 ('DDRS-SEC01', 'Security, Access Control & Incident Response', 'Admin/IT', '2 h', 'ILT', 'F04'),
 ('DDRS-TOT01', 'Train-the-Trainer: Facilitation & Tools', 'Master Trainers', '8 h', 'Workshop', 'Role track'),
 ('DDRS-PUB01', 'Public Services: Verify, Alerts, Grievances', 'Public/Helpdesk', '0.5 h', 'E-learning', 'None'),
]
table(['Code', 'Course', 'Audience', 'Duration', 'Mode', 'Prereq'], [list(c) for c in COURSES], font=8.5)
h(2, '3.3 Role-Based Learning Paths')
PATHS = [
 ('Industry / Applicant', 'F01\u2192F02\u2192F03\u2192F04\u2192IND01\u2192IND02\u2192IND03\u2192IND04 (+IND05/06/07 as applicable)'),
 ('CDSCO Officer (Review)', 'F01-F04\u2192CEN01\u2192CEN02\u2192CEN03\u2192ANA01'),
 ('Drug Inspector', 'F01-F04\u2192CEN01\u2192CEN04\u2192CEN05'),
 ('State Licensing Authority', 'F01-F04\u2192CEN01\u2192CEN02\u2192CEN03\u2192CEN04\u2192CEN05\u2192ANA01\u2192ANA02'),
 ('Laboratory', 'F01-F04\u2192LAB01\u2192LAB02\u2192LAB03'),
 ('Administrator', 'F01-F04\u2192ADM01\u2192ADM02\u2192ADM03\u2192INT01\u2192SEC01'),
 ('Master Trainer', 'Relevant role track + TOT01'),
]
table(['Audience', 'Recommended Path'], [list(p) for p in PATHS], font=9)
pb()

# ============ 4. DELIVERY ============
h(1, '4. Delivery Methods & Materials')
h(2, '4.1 Delivery Channels')
for t in ['Self-paced e-learning (LMS-hosted modules, micro-videos, simulations).',
          'Virtual instructor-led training (VILT) over webinar platform.',
          'On-site instructor-led training at CDSCO HQ, zonal/port offices and State capitals.',
          'Hands-on labs on a dedicated, seeded training (sandbox) environment.',
          'Webinars, workshops and communities of practice for ongoing learning.']:
    bullet(t)
h(2, '4.2 Training Materials')
table(['Artefact', 'Description', 'Format'],
      [['User Manuals (per role)', 'Step-by-step task guides with screenshots', 'PDF/online'],
       ['Quick Reference Guides', 'One-page job aids per key task', 'PDF/printable'],
       ['E-learning Modules', 'Narrated, interactive, assessed', 'SCORM/HTML5'],
       ['Tutorial Videos', 'Short how-to clips per feature', 'MP4'],
       ['Hands-on Lab Guides', 'Guided exercises on the sandbox', 'PDF/online'],
       ['Knowledge Base / FAQs', 'Searchable self-service articles', 'Web'],
       ['Trainer Kits', 'Slides, facilitator notes, exercise data', 'PPTX/PDF'],
       ['Release Notes & Change Bulletins', 'What changed each release', 'Web/email']])
h(2, '4.3 Training (Sandbox) Environment')
para('A dedicated training environment mirrors production and is pre-seeded with '
     'realistic, India-contextualised demonstration data (500+ records per flow). It '
     'allows risk-free practice of every workflow (application lifecycle, inspections, '
     'lab testing, enforcement, vigilance) without affecting live data, and is reset on '
     'a schedule between cohorts.')
pb()

# ============ 5. TRAIN THE TRAINER ============
h(1, '5. Train-the-Trainer (ToT) & Master Trainers')
para('A multiplier model develops a cadre of master trainers who then deliver training '
     'in their States/zones, ensuring scalable, sustainable capacity beyond the SSP team.')
h(2, '5.1 Cadre Structure')
table(['Tier', 'Who', 'Responsibility'],
      [['Lead Trainers (SSP)', 'SSP training specialists', 'Design curricula, train master trainers, quality assurance'],
       ['Master Trainers', 'Nominated CDSCO & State officers', 'Deliver role training in their region; first-line mentoring'],
       ['Departmental Champions', 'Power users per office/lab', 'Peer support, adoption, feedback collection']])
h(2, '5.2 ToT Programme')
for t in ['Selection of master trainers (2\u20134 per State/zone, per portal).',
          'Intensive role-track training + facilitation skills (DDRS-TOT01).',
          'Certification of master trainers via assessment + teach-back.',
          'Trainer kits, sandbox access and ongoing trainer community.']:
    bullet(t)
pb()

# ============ 6. STATE ONBOARDING & ROLLOUT ============
h(1, '6. State Onboarding & Phased Rollout')
para('Onboarding and training are rolled out in waves across all 36 States/UTs. Each '
     'wave includes readiness checks, master-trainer certification, role-based training, '
     'hands-on labs, go-live support and a stabilisation review. State onboarding '
     'progress is a tracked parameter linked to O&M performance.')
h(2, '6.1 Per-State Onboarding Steps')
for t in ['Stakeholder mapping and user/role enrolment.',
          'Master-trainer nomination and ToT certification.',
          'Role-based training (e-learning + VILT + labs).',
          'Data readiness and (where applicable) migration validation.',
          'Go-live with on-site/virtual hypercare support.',
          'Post-go-live review, adoption metrics and refresher scheduling.']:
    num(t)
h(2, '6.2 Indicative Rollout Waves')
table(['Wave', 'Timeframe', 'Coverage', 'Focus'],
      [['Wave 0 \u2013 Pilot', 'Months 1\u20132', 'CDSCO HQ + 2 zonal offices + 1 lab', 'Internal core, ToT seed, feedback'],
       ['Wave 1', 'Months 2\u20134', 'High-volume States (e.g., MH, GJ, TN, KA, TG, UP)', 'SLA, licensing, inspections'],
       ['Wave 2', 'Months 4\u20137', 'Next 10\u201312 States/UTs', 'Full role tracks, labs'],
       ['Wave 3', 'Months 7\u201310', 'Remaining States/UTs', 'Complete national coverage'],
       ['Wave 4 \u2013 Deepening', 'Months 10\u201336', 'All + industry/public at scale', 'Advanced topics, refreshers, new releases']])
h(2, '6.3 Indicative Training Calendar (per wave)')
table(['Week', 'Activity'],
      [['W1', 'Readiness check, enrolment, master-trainer ToT'],
       ['W2', 'Foundation e-learning (all users) + role-track VILT begins'],
       ['W3', 'Hands-on labs (sandbox) per role; assessments'],
       ['W4', 'Go-live + hypercare; daily clinics; feedback capture'],
       ['W5\u20136', 'Stabilisation, refresher clinics, certification close-out']])
pb()

# ============ 7. ASSESSMENT & CERTIFICATION ============
h(1, '7. Assessment, Certification & Competency')
h(2, '7.1 Assessment Methods')
for t in ['Pre-assessment (baseline) and post-assessment (knowledge gain).',
          'Hands-on task-based assessment on the sandbox (can the user do the task?).',
          'Module quizzes within e-learning (auto-graded, pass mark 70%).',
          'Trainer teach-back for master-trainer certification.']:
    bullet(t)
h(2, '7.2 Certification Levels')
table(['Certificate', 'Criteria'],
      [['DDRS Certified User (role)', 'Complete role path + pass assessment (\u226570%)'],
       ['DDRS Certified Power User', 'Advanced track + task assessment'],
       ['DDRS Certified Master Trainer', 'Role track + ToT + teach-back evaluation']])
h(2, '7.3 Competency Framework')
table(['Level', 'Definition'],
      [['Foundation (F)', 'Can navigate and perform basic, guided tasks'],
       ['Intermediate (I)', 'Performs core role tasks independently'],
       ['Advanced (A)', 'Handles exceptions, mentors others, configures where permitted']])
pb()

# ============ 8. EVALUATION & KPIs ============
h(1, '8. Evaluation, KPIs & Continuous Improvement')
h(2, '8.1 Kirkpatrick Evaluation')
table(['Level', 'Measures', 'Method'],
      [['1 Reaction', 'Satisfaction, relevance', 'Post-session survey'],
       ['2 Learning', 'Knowledge/skill gain', 'Pre/post assessment'],
       ['3 Behaviour', 'On-the-job application', 'Usage analytics, supervisor review'],
       ['4 Results', 'SLA, error/ticket reduction, adoption', 'DDRS dashboards & audit']])
h(2, '8.2 Training KPIs')
table(['KPI', 'Target'],
      [['Users trained (per wave)', '\u2265 95% of enrolled'],
       ['Assessment pass rate', '\u2265 85%'],
       ['Master trainers certified', '\u2265 2 per State/portal'],
       ['Post-training satisfaction (CSAT)', '\u2265 4.2 / 5'],
       ['Support tickets per user (30 days post-go-live)', 'Declining trend'],
       ['SLA-compliant processing post-training', 'Improving trend']])
h(2, '8.3 Continuous Improvement')
para('Feedback, assessment data and support trends are reviewed each wave; materials and '
     'sessions are updated accordingly. New releases trigger change bulletins, refreshed '
     'micro-modules and targeted webinars.')
pb()

# ============ 9. HELPDESK & KM ============
h(1, '9. Helpdesk, Knowledge Management & Post-Go-Live Support')
for t in ['IVRS-enabled, multilingual helpdesk with ticketing, escalation and SLAs.',
          'Searchable knowledge base, FAQs and how-to videos (in-product and public).',
          'Hypercare during go-live: daily clinics and rapid-response support.',
          'Communities of practice and periodic webinars for ongoing learning.',
          'Quarterly refreshers and onboarding for new joiners.']:
    bullet(t)
pb()

# ============ 10. GOVERNANCE & RACI ============
h(1, '10. Governance, Roles & Responsibilities')
h(2, '10.1 Training Governance')
para('A Training Steering Group (CDSCO + SSP) owns the plan, approves curricula and '
     'schedules, and reviews KPIs at each wave. A Training PMO coordinates logistics, '
     'enrolment, materials and reporting.')
h(2, '10.2 RACI')
table(['Activity', 'CDSCO', 'SSP Training', 'Master Trainers', 'State PMO'],
      [['Curriculum design', 'A', 'R', 'C', 'I'],
       ['Material development', 'A', 'R', 'C', 'I'],
       ['ToT delivery', 'A', 'R', 'C', 'I'],
       ['Role training delivery', 'I', 'C', 'R', 'A'],
       ['Sandbox environment', 'A', 'R', 'I', 'I'],
       ['Enrolment & logistics', 'C', 'C', 'I', 'R'],
       ['Assessment & certification', 'A', 'R', 'C', 'I'],
       ['Evaluation & reporting', 'A', 'R', 'C', 'C']])
pb()

# ============ 11. RISK ============
h(1, '11. Training Risk Management')
table(['Risk', 'Mitigation'],
      [['Low attendance / adoption', 'Mandated role certification; leadership sponsorship; incentives'],
       ['Geographic dispersion', 'Blended/VILT + master-trainer multiplier model'],
       ['Language barriers', 'Bilingual materials; regional-language support'],
       ['Skill variance', 'Pre-assessment + adaptive paths; foundation track'],
       ['Release churn', 'Change bulletins + micro-module refreshers'],
       ['Trainer attrition', 'Larger master-trainer pool; recorded sessions; trainer kits']])
pb()

# ============ 12. BUDGET (indicative) ============
h(1, '12. Indicative Budget Heads')
para('Training costs are part of the SSP\u2019s operational scope. The indicative heads '
     'below support planning; actual figures are governed by the contract.')
table(['Head', 'Notes'],
      [['Content development', 'E-learning, videos, manuals, trainer kits'],
       ['LMS & sandbox hosting', 'Platform, environment refresh, access management'],
       ['Delivery (VILT/on-site)', 'Trainer effort, travel for on-site waves'],
       ['Master-trainer programme', 'ToT workshops, certification'],
       ['Helpdesk & support', 'IVRS, agents, knowledge base upkeep'],
       ['Evaluation & reporting', 'Assessment tooling, analytics']])
pb()

# ============ APPENDICES ============
h(1, 'Appendix A \u2014 Sample Course Outline (DDRS-CEN02)')
para('Course: Application Review, Queries & Decisions \u2014 audience: CDSCO/State review '
     'officers \u2014 duration: 2.5 h \u2014 mode: ILT + hands-on lab.')
table(['Time', 'Segment', 'Activity'],
      [['00:00', 'Orientation', 'Objectives, the work queue and masked allocation'],
       ['00:20', 'Review workflow', 'Walkthrough of states: pre-screening \u2192 review \u2192 recommend'],
       ['00:50', 'Raising queries', 'Query types, applicant response loop'],
       ['01:20', 'Hands-on lab', 'Process 3 seeded applications end-to-end'],
       ['02:00', 'Decisions & issuance', 'Approve/reject; auto-issuance & QR'],
       ['02:20', 'Assessment', 'Task-based assessment + Q&A']])
h(1, 'Appendix B \u2014 Sample Assessment Items')
for t in ['Demonstrate creating, submitting and paying for an application (task).',
          'Identify the correct action when a query response is received (MCQ).',
          'Locate an entity\u2019s licence and verify it via QR (task).',
          'Explain the masked-allocation rule and manual override (short answer).',
          'Record an inspection finding with correct severity (task).']:
    num(t)
h(1, 'Appendix C \u2014 Attendance & Feedback Templates')
table(['Field', 'Example'],
      [['Course / Code', 'DDRS-CEN02'], ['Date / Wave', '\u2014'], ['Trainer', '\u2014'],
       ['Attendee / Role / Office', '\u2014'], ['Pre-score / Post-score', '\u2014'],
       ['Result (Pass/Fail)', '\u2014'], ['CSAT (1\u20135)', '\u2014'], ['Comments', '\u2014']])
h(1, 'Appendix D \u2014 Glossary')
table(['Term', 'Meaning'],
      [['ADDIE', 'Analyse, Design, Develop, Implement, Evaluate (ID model)'],
       ['VILT', 'Virtual Instructor-Led Training'], ['ILT', 'Instructor-Led Training'],
       ['ToT', 'Train-the-Trainer'], ['LMS', 'Learning Management System'],
       ['CSAT', 'Customer/Trainee Satisfaction score'], ['Sandbox', 'Seeded non-production training environment'],
       ['Hypercare', 'Intensive post-go-live support period'], ['SHRESTH', 'State regulatory performance benchmarking index'],
       ['SLA', 'Service Level Agreement']])

para('\n'); para('\u2014 End of Document \u2014', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, color=NAVY)

OUT = os.path.join(HERE, 'DDRS_Capacity_Building_and_Training_Plan.docx')
doc.save(OUT)
print('SAVED', OUT, '| paragraphs=%d tables=%d' % (len(doc.paragraphs), len(doc.tables)))
